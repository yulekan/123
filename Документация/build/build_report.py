# -*- coding: utf-8 -*-
"""Сборка отчёта по технологической практике на основе шаблона колледжа."""

import os
import shutil
import sys
from copy import deepcopy

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from docx_helpers import Builder, set_pstyle, fmt, run_fmt, TNR, COURIER
from content_diary import DAYS
import content_part1
import content_part2

TEMPLATE = r'C:\Users\gopot\AppData\Local\Temp\otchet.docx'
BASE = r'D:\визуал студио\Учёт спецодежды и СИЗ\Исправленная версия'
DOC_DIR = os.path.join(BASE, 'Документация')
OUT = os.environ.get('PPE_REPORT_OUT') or os.path.join(
    DOC_DIR, 'Отчет_по_технологической_практике_Кишко_ИТ409.docx')
IMAGES = os.path.join(DOC_DIR, 'Диаграммы')
SHOTS = os.path.join(DOC_DIR, 'Скриншоты')
SELFTEST = os.path.join(DOC_DIR, 'selftest.log')

TOPIC_LINES = ('Разработка информационной системы учёта спецодежды и СИЗ',
               'для ООО «Велес-Мит»')

WORKS = ('в изучении организационно-функциональной структуры организации, сбор и анализ требований '
         'к программному средству; в анализе действующего порядка учёта спецодежды и СИЗ, норм выдачи '
         'и сроков носки; в построении функциональной модели IDEF0 и восьми диаграмм UML; '
         'в проектировании и реализации базы данных из девяти отношений средствами SQLite; '
         'в разработке настольного приложения на языке C# (Windows Forms), обеспечивающего учёт прихода, '
         'выдачи и возврата СИЗ с контролем норм и складских остатков; в разработке пяти печатных форм '
         'с предпросмотром и сохранением '
         'в PDF; в тестировании приложения по чек-листу, тест-кейсам и автоматизированному дымовому '
         'тесту; в подготовке руководств пользователя, оператора и программиста')


def norm_ws(text):
    return ' '.join((text or '').split())


def find_paragraph(doc, predicate):
    for par in doc.paragraphs:
        if predicate(norm_ws(par.text)):
            return par
    return None


def fill_topic_line(label_par):
    """Тема индивидуального задания: подпись без двоеточия, значение — на линии.

    Как ФИО на титуле: текст подчёркнут, а линия тянется до правого края
    (табуляция с заливкой «подчёркивание», как в строке «Руководитель
    практики от предприятия»).
    """
    set_text(label_par, 'Тема индивидуального задания')

    next_el = label_par._p.getnext()
    if next_el is None or next_el.tag != qn('w:p'):
        return
    line_par = Paragraph(next_el, label_par._parent)
    if set((line_par.text or '').strip()) != {'_'}:
        return

    for extra in list(line_par.runs):
        extra._r.getparent().remove(extra._r)

    line_par.paragraph_format.tab_stops.add_tab_stop(
        Cm(16.87), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.LINES)

    for index, text in enumerate(TOPIC_LINES):
        if index:
            line_par.add_run().add_break()
        run = line_par.add_run(text)
        run.bold = True
        run.underline = True
        run.font.size = Pt(14)
        run.font.name = TNR
        run._r.get_or_add_rPr().find(qn('w:rFonts')).set(qn('w:cs'), TNR)
        line_par.add_run().add_tab()


def set_text(par, text, template_par=None):
    """Заменяет текст абзаца, сохраняя форматирование первого запуска."""
    runs = par.runs
    if runs:
        runs[0].text = text
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        run = par.add_run(text)
        if template_par is not None and template_par.runs:
            rPr = template_par.runs[0]._r.find(qn('w:rPr'))
            if rPr is not None:
                run._r.insert(0, deepcopy(rPr))


def set_cell(cell, text, template_cell=None):
    """Заменяет текст ячейки таблицы, сохраняя формат."""
    paragraphs = cell.paragraphs
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    par = cell.paragraphs[0]
    runs = par.runs
    if runs:
        runs[0].text = text
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        run = par.add_run(text)
        source = None
        if template_cell is not None and template_cell.paragraphs[0].runs:
            source = template_cell.paragraphs[0].runs[0]
        if source is not None:
            rPr = source._r.find(qn('w:rPr'))
            if rPr is not None:
                run._r.insert(0, deepcopy(rPr))
            else:
                run_fmt(run, size=12)


def append_cell_text(cell, text):
    """Дописывает в ячейку абзац с текстом, копируя формат последнего абзаца."""
    from docx.text.paragraph import Paragraph

    source = cell.paragraphs[-1]._p
    new_p = deepcopy(source)
    runs = new_p.findall(qn('w:r'))
    for extra in runs[1:]:
        new_p.remove(extra)
    if runs:
        first = runs[0]
        for child in list(first):
            if child.tag in (qn('w:t'), qn('w:br'), qn('w:tab')):
                first.remove(child)
        text_el = OxmlElement('w:t')
        text_el.set(qn('xml:space'), 'preserve')
        text_el.text = text
        first.append(text_el)
        source.addnext(new_p)
        return
    source.addnext(new_p)
    run_fmt(Paragraph(new_p, cell).add_run(text), size=12)


def clear_par(par):
    for run in list(par.runs):
        run._r.getparent().remove(run._r)


# --------------------------------------------------------------- шрифты
FONT = 'Times New Roman'
FONT_ATTRS = ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia')
THEME_ATTRS = ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme')
# символьные шрифты маркеров списков менять нельзя — иначе сломаются глифы
KEEP_FONTS = ('Symbol', 'Wingdings', 'Webdings', 'Marlett')


def normalize_fonts(element):
    """Приводит все шрифты в XML-дереве части документа к Times New Roman."""
    changed = 0
    for rFonts in element.iter(qn('w:rFonts')):
        values = [rFonts.get(qn(a)) for a in FONT_ATTRS + THEME_ATTRS]
        if any(v in KEEP_FONTS for v in values if v):
            continue
        for attr in THEME_ATTRS:
            if rFonts.get(qn(attr)) is not None:
                del rFonts.attrib[qn(attr)]
        for attr in FONT_ATTRS:
            if rFonts.get(qn(attr)) != FONT:
                rFonts.set(qn(attr), FONT)
                changed += 1
    return changed


def fix_theme_fonts(path):
    """Меняет шрифты темы документа (word/theme/theme1.xml) на Times New Roman."""
    import re
    import zipfile

    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == 'word/theme/theme1.xml':
                text = data.decode('utf-8')
                text = re.sub(r'<a:latin typeface="[^"]*"',
                              '<a:latin typeface="%s"' % FONT, text)
                data = text.encode('utf-8')
            dst.writestr(item, data)
    os.replace(tmp, path)


def main():
    os.makedirs(DOC_DIR, exist_ok=True)
    shutil.copyfile(TEMPLATE, OUT)
    doc = Document(OUT)
    b = Builder(doc)
    body = doc.element.body

    # ---------------------------------------------------------------- титул
    for par in doc.paragraphs:
        text = norm_ws(par.text)
        if text.startswith('Руководитель практики от организации') and text.endswith('*'):
            set_text(par, 'Руководитель практики от организации*  ___________________')

    # -------------------------------------------------- общие сведения
    topic_par = find_paragraph(doc, lambda t: t == 'Тема индивидуального задания')
    if topic_par is not None:
        fill_topic_line(topic_par)

    # ------------------------------------------------- характеристика
    works_par = find_paragraph(doc, lambda t: t.startswith('с 07.09.2026 по 31.10.2026. выполнял работы'))
    if works_par is not None:
        set_text(works_par, 'с 07.09.2026 по 31.10.2026 выполнял работы: ' + WORKS + '.')

    # ------------------------------------------------- замена подчёркиваний
    hint = find_paragraph(doc, lambda t: t == '(перечислить основные виды работ)')
    if hint is not None:
        b.delete(hint)

    # ----------------------------------------------------------- дневник
    # Таблица дневника в шаблоне заполнена полностью (даты, наименования
    # работ, часы, отметки), а колонка «№ п/п» нумеруется автоматически
    # (список numId=20), поэтому в неё ничего не пишем.

    # ------------------------------------------------------------- отчёт
    # Вводная фраза (наименование работы по программе практики) сохраняется из
    # шаблона. Содержание работы берётся под фактическую практику: где шаблон
    # содержал только образец для другой специальности — он заменяется, где
    # содержания не было вовсе — абзац с текстом дописывается.
    report = doc.tables[3]
    filled = 0
    replaced = 0
    for index, (_date, _hours, _work, text) in enumerate(DAYS):
        cell = report.rows[index + 1].cells[1]
        heading_seen = False
        sample = None
        extra = []
        for par in cell.paragraphs:
            value = norm_ws(par.text)
            if not heading_seen:
                if value:
                    heading_seen = True
                continue
            if value:
                if sample is None:
                    sample = par
                else:
                    extra.append(par)
        if sample is not None:
            set_text(sample, text)
            for par in extra:
                b.delete(par)
            replaced += 1
        else:
            append_cell_text(cell, text)
        filled += 1
    print('заполнено строк отчёта:', filled, 'из', len(DAYS),
          '| заменён образец шаблона в строках:', replaced)

    # ----------------------------------------------------- оглавление
    # В шаблоне уже есть элемент «Оглавление» (sdt с полем TOC), поэтому
    # подсказку «НУМЕРАЦИЮ РАССТАВИТЬ САМОСТОЯТЕЛЬНО…» удаляем,
    # а готовое оглавление обновляет Word.
    toc_note = find_paragraph(doc, lambda t: t.startswith('НУМЕРАЦИЮ РАССТАВИТЬ'))
    if toc_note is not None:
        b.delete(toc_note)

    # --------------------------------------------------- заголовки разделов
    for title in ('ОБЩИЕ СВЕДЕНИЯ', 'ХАРАКТЕРИСТИКА', 'ДНЕВНИК'):
        par = find_paragraph(doc, lambda t, title=title: t == title)
        if par is not None:
            set_pstyle(par, '13')
            pPr = par._p.get_or_add_pPr()
            direction = pPr.find(qn('w:textDirection'))
            if direction is not None:
                pPr.remove(direction)
            ind = pPr.find(qn('w:ind'))
            if ind is not None:
                pPr.remove(ind)
            fmt(par, align='center', indent=None, spacing=1.0, before=0, after=12, keep=True)
            for run in par.runs:
                run_fmt(run, size=16, bold=True)

    # заголовок ОТЧЕТ — абзац перед таблицей отчёта
    report_tbl = report._tbl
    previous = report_tbl.getprevious()
    while previous is not None and previous.tag != qn('w:p'):
        previous = previous.getprevious()
    if previous is not None:
        from docx.text.paragraph import Paragraph
        par = Paragraph(previous, doc)
        clear_par(par)
        set_pstyle(par, '13')
        fmt(par, align='center', indent=None, spacing=1.0, before=0, after=12, keep=True)
        run_fmt(par.add_run('ОТЧЕТ'), size=16, bold=True)

    # --------------------------------------------- удаление инструкций
    from docx.text.paragraph import Paragraph

    start_el = None
    stop_el = None
    for element in list(body):
        if element.tag != qn('w:p'):
            continue
        text = norm_ws(Paragraph(element, doc).text)
        if start_el is None and text.startswith('ПРИЛОЖЕНИЯ ПО ИНДИВИДУАЛЬНОМУ ЗАДАНИЮ'):
            start_el = element
        elif start_el is not None and stop_el is None and text.startswith('АНКЕТА'):
            stop_el = element
            break
    if start_el is not None and stop_el is not None:
        removing = False
        for element in list(body):
            if element is stop_el:
                break
            if element is start_el:
                removing = True
            if removing:
                body.remove(element)

    # ----------------------------------------------- заголовок ПРИЛОЖЕНИЯ
    appendix_heading = find_paragraph(doc, lambda t: t.startswith('ПРИЛОЖЕНИЯ'))
    if appendix_heading is not None:
        clear_par(appendix_heading)
        pPr = appendix_heading._p.get_or_add_pPr()
        direction = pPr.find(qn('w:textDirection'))
        if direction is not None:
            pPr.remove(direction)
        set_pstyle(appendix_heading, '13')
        fmt(appendix_heading, align='center', indent=None, spacing=1.0, before=0, after=12, keep=True)
        appendix_heading.paragraph_format.page_break_before = True
        run_fmt(appendix_heading.add_run('ПРИЛОЖЕНИЯ'), size=16, bold=True)

        # вводный абзац к разделу приложений
        intro_el = OxmlElement('w:p')
        appendix_heading._p.addnext(intro_el)
        intro = Paragraph(intro_el, doc)
        fmt(intro, align='justify', indent=1.25, spacing=1.5, before=0, after=6)
        run_fmt(intro.add_run(
            'В приложениях приведены материалы, полученные в ходе проектирования и '
            'разработки информационной системы учёта спецодежды и СИЗ: описание '
            'предметной области, модели предметной области на языках UML и IDEF0, '
            'постановка задачи, техническое задание, описание структуры базы данных, '
            'руководство пользователя, программа и методика испытаний, сведения о '
            'программной документации и листинги исходного кода.'), size=14)

    # --------------------------------------------------- вставка разделов
    def img(name):
        return os.path.join(IMAGES, name)

    def shot(name):
        return os.path.join(SHOTS, name)

    def project(relative):
        return os.path.join(BASE, relative)

    appendix_anchor = find_paragraph(doc, lambda t: t.startswith('АНКЕТА'))

    elements = b.capture(lambda: content_part1.build(b, img))
    if appendix_anchor is not None:
        b.move_before(elements, appendix_anchor)

    elements = b.capture(lambda: content_part2.build(b, img, shot, project, SELFTEST))
    if appendix_anchor is not None:
        b.move_before(elements, appendix_anchor)

    # ----------------------------------- заключение и источники до приложений
    def build_conclusion():
        content_part2.build_tail(b)

    elements = b.capture(build_conclusion)
    if appendix_heading is not None:
        b.move_before(elements, appendix_heading)

    # ------------------------------------------------------------- шрифты
    # Везде (текст, таблицы, стили, нумерация, колонтитулы) — Times New Roman.
    font_changes = normalize_fonts(doc.element.body)
    font_changes += normalize_fonts(doc.styles.element)
    for part in doc.part.package.iter_parts():
        if part is doc.part or part.partname in ('/word/styles.xml',):
            continue
        element = getattr(part, 'element', None)
        if element is not None:
            font_changes += normalize_fonts(element)
    print('шрифты Times New Roman: обновлено атрибутов:', font_changes)

    # ------------------------------------- автообновление полей при открытии
    settings = doc.settings.element
    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    settings.append(update)

    doc.save(OUT)
    fix_theme_fonts(OUT)
    print('saved:', OUT)
    print('paragraphs:', len(doc.paragraphs), 'tables:', len(doc.tables))


if __name__ == '__main__':
    main()
